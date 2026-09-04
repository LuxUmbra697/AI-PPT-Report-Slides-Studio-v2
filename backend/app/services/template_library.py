"""Reverse-engineer external PPTX files in ``Template/`` into editable design data.

The parser reads the native object tree rather than guessing from a cover image:
every slide records its normalized coordinates, stacking order, groups, text,
fills, pictures, charts and tables.  Qwen-VL adds a semantic description to
the source image assets; DeepSeek then summarizes the complete layout system.
"""

from __future__ import annotations

import base64
import hashlib
import json
import logging
import mimetypes
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree

import httpx
from langchain_core.messages import HumanMessage, SystemMessage
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic import BaseModel, Field

from app.core.config import Settings, get_settings
from app.core.paths import REPO_ROOT, TEMPLATE_ANALYSIS_CACHE, TEMPLATE_DIR
from app.domain.theme import ThemeOverrides, get_theme
from app.llm.client import create_chat_model

logger = logging.getLogger(__name__)

AnalysisStatus = Literal["not_analyzed", "ready", "partial", "unavailable", "failed"]
_ANALYSIS_SCHEMA_VERSION = 2
_HEX_VALUE = re.compile(r"^[0-9A-Fa-f]{6}$")
_QWEN_IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp"}
_MAX_IMAGE_BYTES = 1_500_000
_MAX_TEXT = 120
_MAX_LLM_ELEMENTS = 36


class ExternalTemplateStyleProfile(BaseModel):
    base_theme_id: str
    overrides: ThemeOverrides


class TemplateRect(BaseModel):
    """Coordinates are normalized to the slide (0 to 1), not raw EMU."""

    left: float
    top: float
    width: float
    height: float


class TemplateBackground(BaseModel):
    kind: Literal["none", "solid", "image", "mixed"] = "none"
    colors: list[str] = Field(default_factory=list)
    asset_id: str | None = None


class TemplateElement(BaseModel):
    id: str
    kind: str
    role: str
    name: str
    rect: TemplateRect
    z_index: int
    source_layer: Literal["master", "layout", "slide"] = "slide"
    group_path: list[str] = Field(default_factory=list)
    text_preview: str | None = None
    colors: list[str] = Field(default_factory=list)
    font_size_pt: float | None = None
    asset_id: str | None = None
    asset_description: str | None = None
    is_background: bool = False
    details: dict[str, str | int | float | bool] = Field(default_factory=dict)


class TemplateSlideAnalysis(BaseModel):
    slide_number: int
    layout_name: str | None = None
    title: str | None = None
    background: TemplateBackground
    element_count: int
    element_counts: dict[str, int] = Field(default_factory=dict)
    elements: list[TemplateElement]
    structural_summary: str
    semantic_role: str | None = None
    visual_summary: str | None = None


class TemplateComponent(BaseModel):
    kind: str
    role: str
    occurrences: int
    sample_slides: list[int]
    description: str


class TemplateLayoutRecipe(BaseModel):
    id: str
    name: str
    slide_numbers: list[int]
    description: str


class TemplateInspection(BaseModel):
    """Full per-slide parsing result. This is deliberately not returned by the list API."""

    template_id: str
    canvas: TemplateRect
    visual_analysis_mode: Literal["native-object-tree", "native-object-tree-with-vision"]
    visual_analysis_status: AnalysisStatus = "not_analyzed"
    slides: list[TemplateSlideAnalysis]
    component_catalog: list[TemplateComponent]
    layout_recipes: list[TemplateLayoutRecipe]
    analyzed_at: datetime | None = None


class ExternalTemplatePublic(BaseModel):
    """Lightweight selector summary; request ``inspection`` for detailed object data."""

    id: str
    name: str
    filename: str
    source_path: str
    size_bytes: int
    modified_at: datetime
    slide_count: int
    layout_count: int
    image_count: int
    element_count: int = 0
    background_count: int = 0
    component_count: int = 0
    text_preview: list[str] = Field(default_factory=list)
    palette: list[str] = Field(default_factory=list)
    fonts: list[str] = Field(default_factory=list)
    aspect_ratio: str
    structural_summary: str | None = None
    style_profile: ExternalTemplateStyleProfile
    analysis_status: AnalysisStatus = "not_analyzed"
    style_summary: str | None = None
    image_summary: str | None = None


@dataclass(frozen=True)
class _TemplateSource:
    id: str
    path: Path
    digest: str


@dataclass
class _ImageAsset:
    id: str
    content_type: str
    payload: bytes
    width_px: int | None
    height_px: int | None
    uses: list[int]
    priority: float = 0


@dataclass
class _ParsedTemplate:
    public: ExternalTemplatePublic
    inspection: TemplateInspection
    assets: dict[str, _ImageAsset]


@dataclass(frozen=True)
class ExternalTemplateRenderLayer:
    """A safe visual layer that can be applied below generated editable content."""

    template_id: str
    source_slide_number: int
    elements: tuple[TemplateElement, ...]
    assets: dict[str, tuple[str, bytes]]


def list_external_templates() -> list[ExternalTemplatePublic]:
    """Read local PPTX structure plus cached AI results; never invokes a model."""

    cache = _read_cache()
    result: list[ExternalTemplatePublic] = []
    for source in _iter_sources():
        try:
            parsed = _inspect_template(source)
            result.append(_with_cache(parsed.public, cache.get(source.digest, {})))
        except Exception as error:  # a corrupt deck must not hide other templates
            logger.warning("无法解析外部 PPT 模板 %s: %s", source.path.name, error)
            result.append(_fallback_template(source))
    return result


def get_external_template(template_id: str) -> ExternalTemplatePublic | None:
    return next((item for item in list_external_templates() if item.id == template_id), None)


def get_external_template_inspection(template_id: str) -> TemplateInspection | None:
    """Return complete slide/object data on demand without spending AI credits."""

    source = next((item for item in _iter_sources() if item.id == template_id), None)
    if source is None:
        return None
    try:
        parsed = _inspect_template(source)
    except Exception as error:
        logger.warning("无法读取模板详情 %s: %s", source.path.name, error)
        return None
    return _apply_cached_inspection(parsed.inspection, _read_cache().get(source.digest, {}))


def get_external_template_asset(template_id: str, asset_id: str) -> tuple[str, bytes] | None:
    """Return only an asset belonging to a known Template/ deck, never an arbitrary path."""

    source = next((item for item in _iter_sources() if item.id == template_id), None)
    if source is None:
        return None
    try:
        asset = _inspect_template(source).assets.get(asset_id)
    except Exception as error:
        logger.warning("无法读取模板图片 %s/%s: %s", template_id, asset_id, error)
        return None
    return (asset.content_type, asset.payload) if asset is not None else None


def get_external_template_render_layer(
    template_id: str, slide_index: int
) -> ExternalTemplateRenderLayer | None:
    """Map an output page to the corresponding source page's non-text visual furniture."""

    source = next((item for item in _iter_sources() if item.id == template_id), None)
    if source is None:
        return None
    try:
        parsed = _inspect_template(source)
    except Exception as error:
        logger.warning("无法读取模板视觉层 %s: %s", template_id, error)
        return None
    if not parsed.inspection.slides:
        return None
    # The first generated page uses the source cover. Later pages follow source order and
    # reuse the final source pattern only after the source deck has been exhausted.
    source_position = min(max(slide_index, 0), len(parsed.inspection.slides) - 1)
    source_slide = parsed.inspection.slides[source_position]
    elements = tuple(item for item in source_slide.elements if _is_template_visual(item))
    assets = {
        asset_id: (parsed.assets[asset_id].content_type, parsed.assets[asset_id].payload)
        for asset_id in {item.asset_id for item in elements if item.asset_id}
        if asset_id in parsed.assets
    }
    return ExternalTemplateRenderLayer(
        template_id=template_id,
        source_slide_number=source_slide.slide_number,
        elements=elements,
        assets=assets,
    )


async def refresh_external_templates(*, force: bool = False) -> list[ExternalTemplatePublic]:
    """Explicit opt-in AI analysis. Image and structure failures degrade independently."""

    cache = _read_cache()
    settings = get_settings()
    parsed: dict[str, _ParsedTemplate] = {}
    for source in _iter_sources():
        try:
            parsed[source.id] = _inspect_template(source)
        except Exception as error:
            logger.warning("无法解析外部 PPT 模板 %s: %s", source.path.name, error)

    for source in _iter_sources():
        current = parsed.get(source.id)
        if current is None:
            continue
        existing = cache.get(source.digest, {})
        if (
            not force
            and existing.get("schema_version") == _ANALYSIS_SCHEMA_VERSION
            and existing.get("analysis_status") in {"ready", "partial"}
        ):
            continue
        cache[source.digest] = await _analyze_with_ai(current, settings)
    _write_cache(cache)
    return [
        _with_cache(parsed[source.id].public, cache.get(source.digest, {}))
        if source.id in parsed
        else _fallback_template(source)
        for source in _iter_sources()
    ]


def _iter_sources() -> list[_TemplateSource]:
    if not TEMPLATE_DIR.exists():
        return []
    return [
        _source_for(path)
        for path in sorted(TEMPLATE_DIR.rglob("*.pptx"), key=lambda item: item.name.lower())
        if path.is_file() and not path.name.startswith("~$")
    ]


def _source_for(path: Path) -> _TemplateSource:
    relative = path.relative_to(REPO_ROOT).as_posix()
    return _TemplateSource(
        id="external-" + hashlib.sha1(relative.encode("utf-8")).hexdigest()[:12],
        path=path,
        digest=hashlib.sha256(path.read_bytes()).hexdigest(),
    )


def _inspect_template(source: _TemplateSource) -> _ParsedTemplate:
    return _inspect_template_cached(source.id, str(source.path.resolve()), source.digest)


@lru_cache(maxsize=12)
def _inspect_template_cached(template_id: str, raw_path: str, digest: str) -> _ParsedTemplate:
    return _inspect_template_uncached(
        _TemplateSource(id=template_id, path=Path(raw_path), digest=digest)
    )


def _inspect_template_uncached(source: _TemplateSource) -> _ParsedTemplate:
    presentation = Presentation(source.path)
    canvas_width, canvas_height = presentation.slide_width, presentation.slide_height
    theme_colors: Counter[str] = Counter(_theme_xml_colors(source.path))
    fill_colors: Counter[str] = Counter()
    text_colors: Counter[str] = Counter()
    fonts: Counter[str] = Counter()
    text_preview: list[str] = []
    assets: dict[str, _ImageAsset] = {}
    slides: list[TemplateSlideAnalysis] = []

    for number, slide in enumerate(presentation.slides, start=1):
        _collect_fill(getattr(slide.background, "fill", None), fill_colors)
        elements = _read_slide(
            slide,
            number,
            canvas_width,
            canvas_height,
            assets,
            fonts,
            text_colors,
            text_preview,
        )
        background = _infer_background(slide, elements)
        if background.asset_id:
            for element in elements:
                if element.asset_id == background.asset_id:
                    element.is_background = True
                    if element.asset_id and element.asset_id in assets:
                        assets[element.asset_id].priority += 20
        for element in elements:
            fill_colors.update(element.colors)
        slides.append(
            TemplateSlideAnalysis(
                slide_number=number,
                layout_name=_layout_name(slide),
                title=_title_of(elements),
                background=background,
                element_count=len(elements),
                element_counts=dict(Counter(item.kind for item in elements)),
                elements=elements,
                structural_summary=_slide_summary(elements, background),
            )
        )

    for color, count in fill_colors.items():
        theme_colors[color] += count * 280
    for color, count in text_colors.items():
        theme_colors[color] += count * 3
    for asset in assets.values():
        theme_colors.update(_image_palette(asset.payload))
    global_palette = [color for color, _ in theme_colors.most_common(8)]
    # A red cover background can be outnumbered by white content cards in the whole deck.
    # Prefer the actual cover background for the applied theme so selecting a red template
    # does not incorrectly turn the project beige.
    palette = _profile_palette(slides, assets, global_palette)
    base_theme_id, overrides = _make_style_profile(palette)
    components = _component_catalog(slides)
    recipes = _layout_recipes(slides)
    stat = source.path.stat()
    image_count = sum(item.kind == "picture" for slide in slides for item in slide.elements)
    public = ExternalTemplatePublic(
        id=source.id,
        name=source.path.stem,
        filename=source.path.name,
        source_path=source.path.relative_to(REPO_ROOT).as_posix(),
        size_bytes=stat.st_size,
        modified_at=datetime.fromtimestamp(stat.st_mtime, tz=UTC),
        slide_count=len(slides),
        layout_count=len(presentation.slide_layouts),
        image_count=image_count,
        element_count=sum(slide.element_count for slide in slides),
        background_count=sum(slide.background.kind != "none" for slide in slides),
        component_count=len(components),
        text_preview=_dedupe(text_preview),
        palette=palette,
        fonts=[font for font, _ in fonts.most_common(6)],
        aspect_ratio=f"{canvas_width / canvas_height:.2f}:1" if canvas_height else "16:9",
        structural_summary=(
            f"已读取 {len(slides)} 页、{sum(slide.element_count for slide in slides)} 个原生对象、"
            f"{image_count} 个图片对象和 {len(components)} 类重复组件。"
        ),
        style_profile=ExternalTemplateStyleProfile(
            base_theme_id=base_theme_id,
            overrides=overrides,
        ),
    )
    return _ParsedTemplate(
        public=public,
        inspection=TemplateInspection(
            template_id=source.id,
            canvas=TemplateRect(left=0, top=0, width=1, height=1),
            visual_analysis_mode="native-object-tree",
            slides=slides,
            component_catalog=components,
            layout_recipes=recipes,
        ),
        assets=assets,
    )


def _read_slide(
    slide: Any,
    number: int,
    canvas_width: int,
    canvas_height: int,
    assets: dict[str, _ImageAsset],
    fonts: Counter[str],
    text_colors: Counter[str],
    text_preview: list[str],
) -> list[TemplateElement]:
    elements: list[TemplateElement] = []
    z_index = 0

    def walk(
        shapes: Any, group_path: list[str], source_layer: Literal["master", "layout", "slide"]
    ) -> None:
        nonlocal z_index
        for shape in shapes:
            z_index += 1
            element = _read_shape(
                shape,
                f"s{number}-e{z_index}",
                number,
                z_index,
                group_path,
                canvas_width,
                canvas_height,
                assets,
                fonts,
                text_colors,
                text_preview,
                source_layer,
            )
            elements.append(element)
            if getattr(shape, "shapes", None) is not None:
                walk(shape.shapes, [*group_path, element.id], source_layer)

    # PowerPoint renders slide master first, then slide layout, then slide objects.
    # The first two levels carry fixed backgrounds, title bars and recurring ornaments.
    try:
        walk(slide.slide_layout.slide_master.shapes, [], "master")
        walk(slide.slide_layout.shapes, [], "layout")
    except (AttributeError, TypeError, ValueError):
        logger.debug("页面 %s 没有可读取的母版或版式对象", number)
    walk(slide.shapes, [], "slide")
    return elements


def _read_shape(
    shape: Any,
    element_id: str,
    slide_number: int,
    z_index: int,
    group_path: list[str],
    canvas_width: int,
    canvas_height: int,
    assets: dict[str, _ImageAsset],
    fonts: Counter[str],
    text_colors: Counter[str],
    text_preview: list[str],
    source_layer: Literal["master", "layout", "slide"],
) -> TemplateElement:
    kind = _shape_kind(shape)
    rect = _rect(shape, canvas_width, canvas_height)
    colors: Counter[str] = Counter()
    _collect_fill(getattr(shape, "fill", None), colors)
    _collect_line(getattr(shape, "line", None), colors)
    text, font_size = _text_of(shape, fonts, text_colors, colors)
    if text and len(text_preview) < 12:
        text_preview.append(text)
    asset = _register_picture(shape, assets, slide_number, rect) if kind == "picture" else None
    details: dict[str, str | int | float | bool] = {}
    if kind == "group":
        details["child_count"] = len(shape.shapes)
    if kind in {"chart", "table"}:
        details[kind] = True
    if kind == "shape":
        try:
            auto_shape = shape.auto_shape_type
            if auto_shape is not None:
                details["auto_shape_type"] = str(auto_shape)
        except (AttributeError, TypeError, ValueError):
            pass
    element = TemplateElement(
        id=element_id,
        kind=kind,
        role="其他元素",
        name=str(getattr(shape, "name", "未命名对象")),
        rect=rect,
        z_index=z_index,
        source_layer=source_layer,
        group_path=group_path,
        text_preview=text,
        colors=[color for color, _ in colors.most_common(4)],
        font_size_pt=font_size,
        asset_id=asset.id if asset else None,
        details=details,
    )
    element.role = _role_of(element)
    return element


def _shape_kind(shape: Any) -> str:
    shape_type = getattr(shape, "shape_type", None)
    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if _flag(shape, "has_chart"):
        return "chart"
    if _flag(shape, "has_table"):
        return "table"
    if shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
        return "placeholder"
    connector = getattr(MSO_SHAPE_TYPE, "CONNECTOR", MSO_SHAPE_TYPE.LINE)
    if shape_type in {MSO_SHAPE_TYPE.LINE, connector}:
        return "line"
    if _flag(shape, "has_text_frame"):
        return "text"
    if shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return "shape"
    return "other"


def _flag(shape: Any, name: str) -> bool:
    try:
        return bool(getattr(shape, name, False))
    except (AttributeError, TypeError, ValueError):
        return False


def _rect(shape: Any, width: int, height: int) -> TemplateRect:
    def value(raw: Any, total: int) -> float:
        return round(max(0, min(1, int(raw or 0) / total)), 4) if total else 0

    return TemplateRect(
        left=value(getattr(shape, "left", 0), width),
        top=value(getattr(shape, "top", 0), height),
        width=value(getattr(shape, "width", 0), width),
        height=value(getattr(shape, "height", 0), height),
    )


def _text_of(
    shape: Any, fonts: Counter[str], text_colors: Counter[str], colors: Counter[str]
) -> tuple[str | None, float | None]:
    if not _flag(shape, "has_text_frame"):
        return None, None
    snippets: list[str] = []
    sizes: list[float] = []
    for paragraph in shape.text_frame.paragraphs:
        if text := _compact_text(paragraph.text, limit=_MAX_TEXT):
            snippets.append(text)
        for run in paragraph.runs:
            if run.font.name:
                fonts[run.font.name] += 1
            _collect_font(run.font, text_colors)
            _collect_font(run.font, colors)
            try:
                if run.font.size is not None:
                    sizes.append(float(run.font.size.pt))
            except (AttributeError, TypeError, ValueError):
                pass
    return (
        _compact_text(" / ".join(snippets), limit=_MAX_TEXT) or None,
        max(sizes) if sizes else None,
    )


def _register_picture(
    shape: Any, assets: dict[str, _ImageAsset], slide_number: int, rect: TemplateRect
) -> _ImageAsset | None:
    try:
        image, payload = shape.image, shape.image.blob
        extension = image.ext
    except (AttributeError, TypeError, ValueError):
        return None
    if not payload:
        return None
    asset_id = "asset-" + hashlib.sha1(payload).hexdigest()[:12]
    asset = assets.get(asset_id)
    if asset is None:
        image_width, image_height = _image_size(payload)
        asset = _ImageAsset(
            id=asset_id,
            content_type=mimetypes.guess_type(f"image.{extension}")[0] or "image/png",
            payload=payload,
            width_px=image_width,
            height_px=image_height,
            uses=[],
        )
        assets[asset_id] = asset
    if slide_number not in asset.uses:
        asset.uses.append(slide_number)
    asset.priority += rect.width * rect.height
    return asset


def _image_size(payload: bytes) -> tuple[int | None, int | None]:
    try:
        with Image.open(BytesIO(payload)) as image:
            return image.width, image.height
    except (OSError, ValueError):
        return None, None


def _infer_background(slide: Any, elements: list[TemplateElement]) -> TemplateBackground:
    colors: Counter[str] = Counter()
    _collect_fill(getattr(slide.background, "fill", None), colors)
    covers = [
        item
        for item in elements
        if item.rect.left <= 0.06
        and item.rect.top <= 0.08
        and item.rect.width >= 0.9
        and item.rect.height >= 0.85
    ]
    for item in covers:
        colors.update(item.colors)
    picture = next((item for item in covers if item.kind == "picture" and item.asset_id), None)
    kind: Literal["none", "solid", "image", "mixed"]
    if picture and colors:
        kind = "mixed"
    elif picture:
        kind = "image"
    elif colors:
        kind = "solid"
    else:
        kind = "none"
    return TemplateBackground(
        kind=kind,
        colors=[color for color, _ in colors.most_common(4)],
        asset_id=picture.asset_id if picture else None,
    )


def _layout_name(slide: Any) -> str | None:
    try:
        return str(slide.slide_layout.name) or None
    except (AttributeError, TypeError, ValueError):
        return None


def _title_of(elements: list[TemplateElement]) -> str | None:
    candidates = [
        item
        for item in elements
        if item.kind in {"text", "placeholder"} and item.text_preview and item.rect.top < 0.36
    ]
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda item: (
            item.font_size_pt or 0,
            item.rect.width * item.rect.height,
            -item.rect.top,
        ),
    ).text_preview


def _role_of(element: TemplateElement) -> str:
    rect, area = element.rect, element.rect.width * element.rect.height
    if element.kind == "group":
        return "组合组件"
    if element.kind == "chart":
        return "数据图表"
    if element.kind == "table":
        return "数据表格"
    if element.kind == "line":
        return "连接线或分隔线"
    if element.kind == "picture":
        if rect.width >= 0.9 and rect.height >= 0.85:
            return "全页背景图"
        return "主视觉图片" if area > 0.18 else "图标或装饰" if area < 0.025 else "内容图片"
    if element.kind in {"text", "placeholder"}:
        if rect.top < 0.22 and (element.font_size_pt or 0) >= 18:
            return "页标题"
        return "页眉或标签" if rect.top < 0.2 and rect.height < 0.12 else "正文说明"
    if element.kind == "shape":
        if rect.width >= 0.9 and rect.height >= 0.85:
            return "背景色块"
        return "信息面板" if area > 0.1 else "点缀图案" if area < 0.015 else "卡片或容器"
    return "其他元素"


def _is_template_visual(element: TemplateElement) -> bool:
    """Keep source visual furniture, never overlay source sample text onto new content."""

    if element.kind == "picture":
        # Backgrounds, master assets and source illustrations are all behind generated blocks.
        return element.asset_id is not None
    if element.kind == "line":
        return True
    if element.kind != "shape" or element.text_preview:
        return False
    return element.source_layer != "slide" or element.role in {"背景色块", "点缀图案"}


def _profile_palette(
    slides: list[TemplateSlideAnalysis],
    assets: dict[str, _ImageAsset],
    global_palette: list[str],
) -> list[str]:
    if not slides:
        return global_palette
    cover = slides[0].background
    if cover.asset_id and cover.asset_id in assets:
        cover_colors = [
            color for color, _ in _image_palette(assets[cover.asset_id].payload).most_common(6)
        ]
        if cover_colors:
            return _dedupe([*cover_colors, *global_palette])[:8]
    return _dedupe([*cover.colors, *global_palette])[:8]


def _slide_summary(elements: list[TemplateElement], background: TemplateBackground) -> str:
    roles = Counter(item.role for item in elements if not item.is_background)
    zones: list[str] = []
    if any(item.role == "页标题" for item in elements):
        zones.append("顶部标题区")
    if any(item.role in {"主视觉图片", "内容图片"} for item in elements):
        zones.append("图文内容区")
    if any(item.role in {"信息面板", "卡片或容器"} for item in elements):
        zones.append("卡片/面板组件")
    if any(item.role == "数据图表" for item in elements):
        zones.append("数据可视化区")
    labels = {
        "none": "无显式背景",
        "solid": "纯色背景",
        "image": "全页图片背景",
        "mixed": "图片与色块背景",
    }
    return (
        f"{labels[background.kind]}；{'、'.join(zones) or '自由构图'}；"
        f"主要对象：{'、'.join(roles) or '基础对象'}。"
    )


def _component_catalog(slides: list[TemplateSlideAnalysis]) -> list[TemplateComponent]:
    uses: dict[tuple[str, str], list[int]] = defaultdict(list)
    for slide in slides:
        seen: set[tuple[str, str]] = set()
        for item in slide.elements:
            if item.is_background or item.role in {"其他元素", "正文说明"}:
                continue
            key = (item.kind, item.role)
            if key not in seen:
                uses[key].append(slide.slide_number)
                seen.add(key)
    catalog = [
        TemplateComponent(
            kind=kind,
            role=role,
            occurrences=len(pages),
            sample_slides=pages[:6],
            description=f"出现在 {len(pages)} 页的{role}，可作为重复使用的{kind}组件。",
        )
        for (kind, role), pages in uses.items()
        if len(pages) >= 2
    ]
    return sorted(catalog, key=lambda item: (-item.occurrences, item.role))[:16]


def _layout_recipes(slides: list[TemplateSlideAnalysis]) -> list[TemplateLayoutRecipe]:
    groups: dict[str, list[TemplateSlideAnalysis]] = defaultdict(list)
    for slide in slides:
        counts = slide.element_counts
        fingerprint = ":".join(
            [
                slide.background.kind,
                "title" if slide.title else "untitled",
                "image" if counts.get("picture") else "no-image",
                "chart" if counts.get("chart") else "no-chart",
                "panel"
                if any(item.role in {"信息面板", "卡片或容器"} for item in slide.elements)
                else "free",
            ]
        )
        groups[fingerprint].append(slide)
    recipes: list[TemplateLayoutRecipe] = []
    for index, (fingerprint, grouped) in enumerate(
        sorted(groups.items(), key=lambda item: (-len(item[1]), item[0])), start=1
    ):
        first = grouped[0]
        if first.slide_number == 1:
            name = "封面主视觉"
        elif first.element_counts.get("chart"):
            name = "数据图表页"
        elif any(item.role in {"信息面板", "卡片或容器"} for item in first.elements):
            name = "卡片信息页"
        elif first.background.kind in {"image", "mixed"}:
            name = "图文叙事页"
        else:
            name = "标题正文页"
        recipes.append(
            TemplateLayoutRecipe(
                id=f"layout-{index}",
                name=name,
                slide_numbers=[slide.slide_number for slide in grouped],
                description=f"{first.structural_summary} 结构指纹：{fingerprint}。",
            )
        )
    return recipes[:12]


def _fallback_template(source: _TemplateSource) -> ExternalTemplatePublic:
    stat = source.path.stat()
    return ExternalTemplatePublic(
        id=source.id,
        name=source.path.stem,
        filename=source.path.name,
        source_path=source.path.relative_to(REPO_ROOT).as_posix(),
        size_bytes=stat.st_size,
        modified_at=datetime.fromtimestamp(stat.st_mtime, tz=UTC),
        slide_count=0,
        layout_count=0,
        image_count=0,
        aspect_ratio="16:9",
        style_profile=ExternalTemplateStyleProfile(
            base_theme_id="ivory", overrides=ThemeOverrides()
        ),
        analysis_status="failed",
        style_summary="模板文件存在，但当前无法读取其内部对象结构。",
    )


def _collect_fill(fill: Any, target: Counter[str]) -> None:
    try:
        _append_color(fill.fore_color, target)
    except (AttributeError, TypeError, ValueError):
        pass


def _collect_line(line: Any, target: Counter[str]) -> None:
    try:
        _append_color(line.color, target)
    except (AttributeError, TypeError, ValueError):
        pass


def _collect_font(font: Any, target: Counter[str]) -> None:
    try:
        _append_color(font.color, target)
    except (AttributeError, TypeError, ValueError):
        pass


def _append_color(color: Any, target: Counter[str]) -> None:
    value = str(getattr(color, "rgb", ""))
    if _HEX_VALUE.match(value):
        target[f"#{value.upper()}"] += 1


def _theme_xml_colors(path: Path) -> list[str]:
    result: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("ppt/theme/") and name.endswith(".xml"):
                for node in ElementTree.fromstring(archive.read(name)).iter():
                    if _HEX_VALUE.match(value := node.attrib.get("val", "")):
                        result.append(f"#{value.upper()}")
    return result


def _image_palette(payload: bytes) -> Counter[str]:
    try:
        with Image.open(BytesIO(payload)) as image:
            image = image.convert("RGB")
            image.thumbnail((120, 120))
            pixels = image.get_flattened_data()
    except (OSError, ValueError):
        return Counter()
    return Counter(
        "#" + "".join(f"{channel // 32 * 32:02X}" for channel in pixel) for pixel in pixels
    )


def _make_style_profile(palette: list[str]) -> tuple[str, ThemeOverrides]:
    fallback = get_theme("ivory")
    background = palette[0] if palette else fallback.palette.background
    dark = _luminance(background) < 0.42
    base_theme_id = "midnight" if dark else "ivory"
    base = get_theme(base_theme_id)
    accent = _pick_accent(palette, background) or base.palette.accent
    ink = "#F6F8FF" if dark else "#1A1917"
    return (
        base_theme_id,
        ThemeOverrides.model_validate(
            {
                "palette": {
                    "background": background,
                    "surface": _mix(background, "#FFFFFF", 0.09 if dark else 0.64),
                    "ink": ink,
                    "ink_soft": "#C9D1E3" if dark else "#4A463E",
                    "ink_muted": "#8F9AB3" if dark else "#8A8478",
                    "accent": accent,
                    "accent_soft": _mix(background, accent, 0.18),
                    "line": _mix(background, ink, 0.16),
                    "line_strong": _mix(background, ink, 0.30),
                },
                "shape": {"radius_pt": 8 if dark else 4, "bullet_marker": "dot"},
            }
        ),
    )


def _luminance(value: str) -> float:
    red, green, blue = _rgb(value)
    return (0.2126 * red + 0.7152 * green + 0.0722 * blue) / 255


def _saturation(value: str) -> float:
    red, green, blue = _rgb(value)
    maximum = max(red, green, blue)
    return 0 if maximum == 0 else (maximum - min(red, green, blue)) / maximum


def _pick_accent(palette: list[str], background: str) -> str | None:
    candidates = [color for color in palette if color != background]
    if not candidates:
        return None
    base_luminance = _luminance(background)
    return max(
        candidates,
        key=lambda color: _saturation(color) * 1.1
        + abs(_luminance(color) - base_luminance) * 0.4
        + (_luminance(color) if base_luminance < 0.5 else 1 - _luminance(color)) * 0.7,
    )


def _mix(left: str, right: str, ratio: float) -> str:
    values = [round(a + (b - a) * ratio) for a, b in zip(_rgb(left), _rgb(right), strict=True)]
    return "#" + "".join(f"{value:02X}" for value in values)


def _rgb(value: str) -> tuple[int, int, int]:
    return int(value[1:3], 16), int(value[3:5], 16), int(value[5:7], 16)


async def _analyze_with_ai(parsed: _ParsedTemplate, settings: Settings) -> dict[str, Any]:
    image_summary, descriptions = await _qwen_assets(parsed.assets, settings)
    style_summary, notes = await _deepseek_structure(
        parsed.public, parsed.inspection, descriptions, settings
    )
    configured = bool(settings.llm_api_key.strip() or _vision_key(settings))
    state: AnalysisStatus = (
        "ready"
        if style_summary and image_summary
        else "partial"
        if style_summary or image_summary
        else "failed"
        if configured
        else "unavailable"
    )
    return {
        "schema_version": _ANALYSIS_SCHEMA_VERSION,
        "analysis_status": state,
        "style_summary": style_summary,
        "image_summary": image_summary,
        "asset_descriptions": descriptions,
        "slide_notes": notes,
        "analyzed_at": datetime.now(UTC).isoformat(),
    }


async def _qwen_assets(
    assets: dict[str, _ImageAsset], settings: Settings
) -> tuple[str | None, dict[str, str]]:
    key = _vision_key(settings)
    if not key:
        return None, {}
    candidates = [
        asset
        for asset in sorted(assets.values(), key=lambda item: (-item.priority, item.id))
        if asset.content_type in _QWEN_IMAGE_TYPES and 0 < len(asset.payload) <= _MAX_IMAGE_BYTES
    ][: settings.template_analysis_max_images]
    descriptions: dict[str, str] = {}
    summaries: list[str] = []
    for offset in range(0, len(candidates), 3):
        batch = candidates[offset : offset + 3]
        content: list[dict[str, Any]] = [
            {
                "type": "text",
                "text": (
                    "这些是用户拥有的 PPT 模板图片资产。不要识别个人身份。"
                    "仅返回 JSON：{\"assets\":[{\"id\":\"asset-id\",\"description\":\"中文短语\"}],"
                    "\"summary\":\"中文短句\"}。"
                ),
            }
        ]
        for asset in batch:
            content.extend(
                [
                    {"type": "text", "text": f"资产 id：{asset.id}"},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:{asset.content_type};base64,"
                            f"{base64.b64encode(asset.payload).decode('ascii')}"
                        },
                    },
                ]
            )
        try:
            raw = await _qwen_completion(content, settings, key)
        except (httpx.HTTPError, ValueError, TypeError) as error:
            logger.warning("Qwen-VL 模板资产分析失败：%s", error)
            continue
        result = _json_object(raw)
        if result is None:
            summaries.append(_compact_text(raw, limit=140))
            continue
        if isinstance(result.get("summary"), str):
            summaries.append(_compact_text(result["summary"], limit=140))
        valid_ids = {asset.id for asset in batch}
        for item in result.get("assets", []):
            if (
                isinstance(item, dict)
                and item.get("id") in valid_ids
                and isinstance(item.get("description"), str)
            ):
                descriptions[item["id"]] = _compact_text(item["description"], limit=120)
    summary = _compact_text("；".join(_dedupe(summaries)), limit=220)
    if not summary and descriptions:
        summary = "、".join(_dedupe(list(descriptions.values()))[:5])
    return summary or None, descriptions


async def _qwen_completion(content: list[dict[str, Any]], settings: Settings, key: str) -> str:
    async with httpx.AsyncClient(timeout=settings.image_timeout_seconds) as client:
        response = await client.post(
            f"{settings.template_vision_base_url.rstrip('/')}/chat/completions",
            headers={"Authorization": f"Bearer {key}"},
            json={
                "model": settings.template_vision_model,
                "messages": [{"role": "user", "content": content}],
                "temperature": 0.15,
            },
        )
        response.raise_for_status()
        data = response.json()
    choices = data.get("choices") or []
    return _message_text(choices[0].get("message", {}).get("content", "")) if choices else ""


async def _deepseek_structure(
    template: ExternalTemplatePublic,
    inspection: TemplateInspection,
    descriptions: dict[str, str],
    settings: Settings,
) -> tuple[str | None, dict[str, dict[str, str]]]:
    if not settings.llm_api_key.strip():
        return None, {}
    payload: list[dict[str, Any]] = []
    for slide in inspection.slides:
        major = sorted(
            slide.elements,
            key=lambda item: (
                item.is_background,
                -(item.rect.width * item.rect.height),
                item.z_index,
            ),
        )[:_MAX_LLM_ELEMENTS]
        payload.append(
            {
                "slide_number": slide.slide_number,
                "layout": slide.layout_name,
                "background": slide.background.model_dump(),
                "title": slide.title,
                "elements": [
                    {
                        "kind": item.kind,
                        "role": item.role,
                        "rect": item.rect.model_dump(),
                        "text": item.text_preview,
                        "image": descriptions.get(item.asset_id or ""),
                    }
                    for item in major
                ],
            }
        )
    try:
        message = await create_chat_model(settings).ainvoke(
            [
                SystemMessage(
                    content=(
                        "你是严谨的 PPT 模板逆向分析师。根据原生对象树归纳构图、背景、"
                        "图片用途、排版、组件与文案语气；不可编造。只返回 JSON："
                        '{"style_summary":"不超过120字","slides":[{"slide_number":1,'
                        '"semantic_role":"封面/目录/正文/章节/结尾等","visual_summary":"不超过80字"}]}。'
                    )
                ),
                HumanMessage(
                    content=json.dumps(
                        {
                            "template": template.name,
                            "palette": template.palette,
                            "fonts": template.fonts,
                            "slides": payload,
                        },
                        ensure_ascii=False,
                    )
                ),
            ]
        )
    except Exception as error:
        logger.warning("DeepSeek 模板结构分析失败：%s", error)
        return None, {}
    result = _json_object(_message_text(message.content))
    if result is None:
        return _compact_text(_message_text(message.content), limit=160) or None, {}
    notes: dict[str, dict[str, str]] = {}
    for item in result.get("slides", []):
        if not isinstance(item, dict) or not isinstance(item.get("slide_number"), int):
            continue
        note = {
            key: _compact_text(value, limit=100)
            for key in ("semantic_role", "visual_summary")
            if isinstance(value := item.get(key), str) and value.strip()
        }
        if note:
            notes[str(item["slide_number"])] = note
    summary = result.get("style_summary")
    return _compact_text(summary, limit=180) if isinstance(summary, str) else None, notes


def _vision_key(settings: Settings) -> str:
    if settings.template_vision_api_key.strip():
        return settings.template_vision_api_key.strip()
    return settings.image_api_key.strip() if settings.image_provider == "bailian" else ""


def _with_cache(template: ExternalTemplatePublic, cache: dict[str, Any]) -> ExternalTemplatePublic:
    if cache.get("schema_version") != _ANALYSIS_SCHEMA_VERSION:
        return template
    keys = {"analysis_status", "style_summary", "image_summary"}
    return template.model_copy(update={key: cache[key] for key in keys if key in cache})


def _apply_cached_inspection(
    inspection: TemplateInspection, cache: dict[str, Any]
) -> TemplateInspection:
    if cache.get("schema_version") != _ANALYSIS_SCHEMA_VERSION:
        return inspection
    descriptions = cache.get("asset_descriptions", {})
    notes = cache.get("slide_notes", {})
    slides = []
    for slide in inspection.slides:
        elements = [
            item.model_copy(update={"asset_description": descriptions[item.asset_id]})
            if item.asset_id in descriptions
            else item
            for item in slide.elements
        ]
        note = notes.get(str(slide.slide_number), {})
        slides.append(
            slide.model_copy(
                update={
                    "elements": elements,
                    "semantic_role": note.get("semantic_role"),
                    "visual_summary": note.get("visual_summary"),
                }
            )
        )
    analyzed_at = None
    try:
        analyzed_at = datetime.fromisoformat(cache["analyzed_at"])
    except (KeyError, TypeError, ValueError):
        pass
    return inspection.model_copy(
        update={
            "slides": slides,
            "visual_analysis_mode": "native-object-tree-with-vision",
            "visual_analysis_status": cache.get("analysis_status", "not_analyzed"),
            "analyzed_at": analyzed_at,
        }
    )


def _json_object(value: str) -> dict[str, Any] | None:
    raw = value.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        parsed = json.loads(raw)
    except ValueError:
        start, end = raw.find("{"), raw.rfind("}")
        if start < 0 or end <= start:
            return None
        try:
            parsed = json.loads(raw[start : end + 1])
        except ValueError:
            return None
    return parsed if isinstance(parsed, dict) else None


def _read_cache() -> dict[str, dict[str, Any]]:
    try:
        data = json.loads(TEMPLATE_ANALYSIS_CACHE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def _write_cache(cache: dict[str, dict[str, Any]]) -> None:
    try:
        TEMPLATE_ANALYSIS_CACHE.parent.mkdir(parents=True, exist_ok=True)
        TEMPLATE_ANALYSIS_CACHE.write_text(
            json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except OSError as error:
        logger.warning("模板 AI 分析缓存无法写入：%s", error)


def _message_text(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return " ".join(
            item.get("text", "") if isinstance(item, dict) else str(item) for item in value
        )
    return str(value)


def _compact_text(value: str, *, limit: int = 64) -> str:
    return " ".join(value.split())[:limit]


def _dedupe(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))
