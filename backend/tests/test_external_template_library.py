import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.domain.sample import load_sample_deck
from app.domain.theme import get_theme
from app.render.pptx import render_deck_to_pptx
from app.services.template_library import (
    _QWEN_IMAGE_TYPES,
    _make_style_profile,
    get_external_template_render_layer,
    list_external_templates,
)


def test_external_palette_maps_to_a_valid_editable_theme_profile() -> None:
    base_theme_id, overrides = _make_style_profile(["#FAEFE2", "#A00000", "#E74100"])

    assert get_theme(base_theme_id).id == base_theme_id
    assert overrides.palette is not None
    assert overrides.palette.background == "#FAEFE2"
    assert overrides.palette.accent == "#A00000"


def test_dark_external_palette_prefers_a_bright_accent() -> None:
    base_theme_id, overrides = _make_style_profile(["#000000", "#2080E0", "#002060"])

    assert base_theme_id == "midnight"
    assert overrides.palette is not None
    assert overrides.palette.accent == "#2080E0"


def test_qwen_image_filter_excludes_vector_formats() -> None:
    assert "image/svg+xml" not in _QWEN_IMAGE_TYPES
    assert {"image/jpeg", "image/png"}.issubset(_QWEN_IMAGE_TYPES)


def test_external_template_visual_layer_is_written_to_exported_pptx() -> None:
    templates = list_external_templates()
    if not templates:
        pytest.skip("Template 文件夹中没有可供集成测试的 PPTX")

    layer = get_external_template_render_layer(templates[0].id, slide_index=0)
    assert layer is not None
    assert layer.elements
    assert layer.assets

    output = render_deck_to_pptx(
        load_sample_deck(),
        theme_id=templates[0].style_profile.base_theme_id,
        external_template_id=templates[0].id,
    )
    presentation = Presentation(output)
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in presentation.slides[0].shapes
    )
